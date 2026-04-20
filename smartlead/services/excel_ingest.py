# smartlead/services/excel_ingest.py
import io
import re
from typing import Any

from fastapi import UploadFile
from openpyxl import load_workbook

from smartlead.api.v1.schemas.leads import LeadInput

def _norm_header(h: str) -> str:
    key = (h or "").strip().lower()
    key = re.sub(r"[\s\-]+", "_", key)
    return key

def _cell_to_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()

def parse_leads_xlsx_bytes(raw_bytes: bytes) -> tuple[list[LeadInput], list[str]]:
    errors: list[str] = []
    try:
        wb = load_workbook(filename=io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception as exc:
        return [], [f"Invalid .xlsx file ({exc})."]

    ws = wb.worksheets[0] if wb.worksheets else None
    if ws is None:
        return [], ["Excel file has no worksheets."]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], ["Excel sheet is empty."]

    headers = [ _cell_to_str(h) for h in header_row ]
    if not any(headers):
        return [], ["Excel header row is empty."]

    fieldmap = {_norm_header(h): i for i, h in enumerate(headers) if h}

    def pick(*candidates: str) -> int | None:
        for c in candidates:
            k = _norm_header(c)
            if k in fieldmap:
                return fieldmap[k]
        return None

    idx_company = pick("company_name", "company", "name", "organization")
    idx_website = pick("website", "url", "domain", "company_website")
    idx_email   = pick("contact_email", "email", "e_mail", "contact")

    if idx_company is None:
        return [], [
            "Missing required column. Provide at least one of: "
            "company_name, company, name, organization.",
        ]

    leads: list[LeadInput] = []
    # Excel row numbers are 1-based; data starts after header => row 2
    excel_row_num = 1
    for excel_row_num, row in enumerate(rows_iter, start=2):
        def get_at(idx: int | None) -> str:
            if idx is None:
                return ""
            if idx >= len(row):
                return ""
            return _cell_to_str(row[idx])

        company = get_at(idx_company)
        if not company:
            errors.append(f"Row {excel_row_num}: empty company; skipped.")
            continue

        website = get_at(idx_website)
        email   = get_at(idx_email)

        try:
            leads.append(LeadInput(company_name=company, website=website, contact_email=email))
        except Exception as exc:
            errors.append(f"Row {excel_row_num}: invalid data ({exc}).")

    if not leads and not errors:
        errors.append("No data rows found after the header.")

    return leads, errors

def parse_leads_xls_bytes(raw_bytes: bytes) -> tuple[list[LeadInput], list[str]]:
    # Requires xlrd
    import xlrd  # noqa: PLC0415

    errors: list[str] = []
    try:
        book = xlrd.open_workbook(file_contents=raw_bytes)
    except Exception as exc:
        return [], [f"Invalid .xls file ({exc})."]

    if book.nsheets < 1:
        return [], ["Excel file has no worksheets."]

    sheet = book.sheet_by_index(0)
    if sheet.nrows < 1:
        return [], ["Excel sheet is empty."]

    headers = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
    if not any(headers):
        return [], ["Excel header row is empty."]

    fieldmap = {_norm_header(h): c for c, h in enumerate(headers) if h}

    def pick(*candidates: str) -> int | None:
        for c in candidates:
            k = _norm_header(c)
            if k in fieldmap:
                return fieldmap[k]
        return None

    idx_company = pick("company_name", "company", "name", "organization")
    idx_website = pick("website", "url", "domain", "company_website")
    idx_email   = pick("contact_email", "email", "e_mail", "contact")

    if idx_company is None:
        return [], [
            "Missing required column. Provide at least one of: "
            "company_name, company, name, organization.",
        ]

    def get_at(row_idx: int, col_idx: int | None) -> str:
        if col_idx is None:
            return ""
        if col_idx >= sheet.ncols:
            return ""
        return _cell_to_str(sheet.cell_value(row_idx, col_idx))

    leads: list[LeadInput] = []

    # Data starts from row 1 (0-based index), Excel row number = r + 1
    for r in range(1, sheet.nrows):
        excel_row_num = r + 1

        company = get_at(r, idx_company)
        if not company:
            errors.append(f"Row {excel_row_num}: empty company; skipped.")
            continue

        website = get_at(r, idx_website)
        email   = get_at(r, idx_email)

        try:
            leads.append(
                LeadInput(
                    company_name=company,
                    website=website,
                    contact_email=email,
                )
            )
        except Exception as exc:
            errors.append(f"Row {excel_row_num}: invalid data ({exc}).")

    if not leads and not errors:
        errors.append("No data rows found after the header.")

    return leads, errors

async def parse_leads_excel(file: UploadFile) -> tuple[list[LeadInput], list[str]]:
    raw_bytes = await file.read()
    name = (file.filename or "").lower()

    if name.endswith(".xlsx"):
        return parse_leads_xlsx_bytes(raw_bytes)
    if name.endswith(".xls"):
        return parse_leads_xls_bytes(raw_bytes)

    return [], ["Unsupported file type. Upload .csv, .xlsx, or .xls."]